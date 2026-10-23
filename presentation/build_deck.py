#!/usr/bin/env python3
"""Build the Sudan JJA 2025 presentation from committed analysis outputs.

Run from the repository root with:

    uv run --with python-pptx python presentation/build_deck.py

The deck is written to ``presentation/JJA2025_Sudan_forecast_risk.pptx``.
All source figures and CSV tables are produced by the committed scripts in
``scripts/``; the two summary charts on the Part A and Part B slides are
generated here from those CSVs.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentation"
OUT_FILE = OUT_DIR / "JJA2025_Sudan_forecast_risk.pptx"
TOTAL_SLIDES = 12

# ---------------------------------------------------------------------------
# Palette and type
# ---------------------------------------------------------------------------
NAVY = "#0B1F33"
NAVY_2 = "#15344D"
INK = "#182A3A"
AMBER = "#D9901A"
AMBER_TEXT = "#A96500"
AMBER_LIGHT = "#FBF0DB"
TEAL = "#0F766E"
TEAL_LIGHT = "#E7F3F1"
ORANGE = "#C2410C"
RED = "#B42318"
PAPER = "#FAF7F0"
CARD = "#FFFFFF"
LINE = "#DDD6C8"
MUTED = "#5B6B7C"
LIGHT_MUTED = "#8A98A5"
DARK_MUTED = "#AAB8C4"
WHITE = "#FFFFFF"

FONT = "Arial"
MONO = "Courier New"


def c(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", ""))


# ---------------------------------------------------------------------------
# Small drawing helpers
# ---------------------------------------------------------------------------
def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=13,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    italic=False,
    line_spacing=1.0,
    margin=0.0,
    link=None,
):
    """Add a simple text box. Newlines become separate paragraphs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor

    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = c(color)
        if link and i == 0:
            run.hyperlink.address = link
    return box


def add_rect(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill=CARD,
    line=None,
    line_width=0.75,
    radius=0.06,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shp = slide.shapes.add_shape(
        shape, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    # Remove the theme style reference. LibreOffice otherwise keeps the theme's
    # effect reference and draws a drop shadow even after shadow.inherit = False.
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = c(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = c(line)
        shp.line.width = Pt(line_width)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def add_line(slide, x1, y1, x2, y2, color=LINE, width=0.75):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    style = line._element.find(qn("p:style"))
    if style is not None:
        line._element.remove(style)
    line.line.color.rgb = c(color)
    line.line.width = Pt(width)
    return line


def add_bullets(
    slide,
    x,
    y,
    w,
    h,
    items,
    *,
    size=12,
    color=INK,
    bullet_color=AMBER,
    font=FONT,
    space_after=5,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.20))))
        pPr.set("indent", str(-int(Inches(0.20))))

        bullet = p.add_run()
        bullet.text = "•  "
        bullet.font.name = font
        bullet.font.size = Pt(size)
        bullet.font.bold = True
        bullet.font.color.rgb = c(bullet_color)

        if isinstance(item, tuple):
            lead, tail = item
            r1 = p.add_run()
            r1.text = lead
            r1.font.name = font
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = c(color)
            r2 = p.add_run()
            r2.text = tail
            r2.font.name = font
            r2.font.size = Pt(size)
            r2.font.color.rgb = c(color)
        else:
            r = p.add_run()
            r.text = item
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = c(color)
    return box


def add_header(slide, section, title, *, title_size=27, dark=False):
    label_color = AMBER if dark else AMBER_TEXT
    title_color = WHITE if dark else NAVY
    line_color = "#2A4760" if dark else LINE

    add_rect(
        slide,
        0.65,
        0.52,
        0.10,
        0.10,
        fill=AMBER,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_text(
        slide,
        0.88,
        0.45,
        6.5,
        0.25,
        section.upper(),
        size=9.5,
        color=label_color,
        bold=True,
    )
    add_text(
        slide,
        0.65,
        0.76,
        12.0,
        0.70,
        title,
        size=title_size,
        color=title_color,
        bold=True,
        line_spacing=0.95,
    )
    add_line(slide, 0.65, 1.43, 12.68, 1.43, color=line_color, width=0.75)


def add_footer(slide, number, *, dark=False, total=TOTAL_SLIDES):
    line_color = "#2A4760" if dark else LINE
    text_color = DARK_MUTED if dark else LIGHT_MUTED
    add_line(slide, 0.65, 7.04, 12.68, 7.04, color=line_color, width=0.75)
    add_text(
        slide,
        0.65,
        7.10,
        7.6,
        0.20,
        "Sudan JJA 2025 · local rainfall extremes",
        size=8.5,
        color=text_color,
    )
    add_text(
        slide,
        11.7,
        7.10,
        0.98,
        0.20,
        f"{number:02d} / {total:02d}",
        size=8.5,
        color=text_color,
        align=PP_ALIGN.RIGHT,
    )


def add_caption(slide, x, y, w, text, *, dark=False, size=8.5):
    add_text(
        slide,
        x,
        y,
        w,
        0.38,
        text,
        size=size,
        color=DARK_MUTED if dark else MUTED,
        line_spacing=1.0,
    )


def add_picture_fit(
    slide,
    path,
    x,
    y,
    max_w,
    max_h,
    *,
    border=True,
    alt=None,
    border_color=LINE,
):
    """Place an image inside a box, preserving its aspect ratio."""
    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = iw / ih
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    left = x + (max_w - w) / 2
    top = y + (max_h - h) / 2
    pic = slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), width=Inches(w), height=Inches(h)
    )
    if alt:
        pic._element.nvPicPr.cNvPr.set("descr", alt)
    if border:
        add_rect(
            slide,
            left,
            top,
            w,
            h,
            fill=None,
            line=border_color,
            line_width=0.75,
            radius=0.0,
            shape=MSO_SHAPE.RECTANGLE,
        )
    return left, top, w, h


def add_chip(slide, x, y, w, h, text, *, fill=NAVY, color=WHITE, size=10):
    chip = add_rect(slide, x, y, w, h, fill=fill, line=None, radius=0.18)
    add_text(
        slide,
        x,
        y + 0.02,
        w,
        h - 0.02,
        text,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return chip


def add_link_overlay(slide, x, y, w, h, url):
    """Invisible click target so linked text keeps the deck's own colours."""
    overlay = add_rect(
        slide,
        x,
        y,
        w,
        h,
        fill=None,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    overlay.click_action.hyperlink.address = url
    return overlay


def slide_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------------------
# Summary charts from committed CSVs
# ---------------------------------------------------------------------------
def make_part_a_chart(path: Path) -> None:
    df = pd.read_csv(ROOT / "out/rainfall_thresholds/thresholds.csv")
    d = (
        df[df["duration"].eq("24hr")]
        .pivot(index="town", columns="return_period_years", values="threshold_mm")
        .sort_values(2)
    )
    order = list(d.index)
    y = np.arange(len(order))
    height = 0.23
    colors = {2: AMBER, 10: TEAL, 50: NAVY}
    labels = {2: "2-year", 10: "10-year", 50: "50-year"}

    plt.rcParams.update(
        {
            "font.family": "Liberation Sans",
            "font.size": 9,
            "axes.edgecolor": LINE,
            "axes.labelcolor": INK,
            "text.color": INK,
            "xtick.color": MUTED,
            "ytick.color": INK,
        }
    )
    fig, ax = plt.subplots(figsize=(8.6, 4.5), dpi=220)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    for i, rp in enumerate([2, 10, 50]):
        vals = d.loc[order, rp].to_numpy(dtype=float)
        bars = ax.barh(
            y + (i - 1) * height,
            vals,
            height=height,
            color=colors[rp],
            edgecolor="none",
            label=labels[rp],
            zorder=3,
        )
        for bar, val in zip(bars, vals):
            ax.text(
                val + 1.2,
                bar.get_y() + bar.get_height() / 2,
                f"{val:.1f}",
                va="center",
                ha="left",
                fontsize=7.6,
                color=INK,
            )

    ax.set_yticks(y)
    ax.set_yticklabels(order, fontsize=10, fontweight="bold")
    ax.invert_yaxis()
    ax.set_xlim(0, 104)
    ax.set_xlabel("24-hour rainfall (mm)", fontsize=10, labelpad=7)
    ax.xaxis.grid(True, color="#ECE7DC", linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(LINE)
    ax.spines["bottom"].set_color(LINE)
    ax.tick_params(axis="both", length=0, pad=5)
    ax.legend(
        loc="upper center",
        bbox_to_anchor=(0.5, 1.12),
        ncol=3,
        frameon=False,
        fontsize=9,
        handlelength=1.2,
        columnspacing=1.5,
    )
    ax.set_title(
        "Gumbel fit to 27 annual maxima (1998–2024)",
        loc="left",
        fontsize=9,
        color=MUTED,
        pad=9,
    )
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_part_b_chart(path: Path) -> None:
    df = pd.read_csv(
        ROOT / "out/forecast_thresholds/07_forecast_threshold_counts.csv"
    )
    label_map = {
        "fixed 20 mm/day": "Fixed 20 mm",
        "CMORPH 2-year": "2-year level",
        "CMORPH 10-year": "10-year level",
    }
    d = df.pivot(index="town", columns="threshold_label", values="member_days_above")
    order = list(d["fixed 20 mm/day"].sort_values(ascending=False).index)
    series = [
        ("fixed 20 mm/day", AMBER),
        ("CMORPH 2-year", TEAL),
        ("CMORPH 10-year", NAVY),
    ]

    plt.rcParams.update(
        {
            "font.family": "Liberation Sans",
            "font.size": 9,
            "axes.edgecolor": LINE,
            "axes.labelcolor": INK,
            "text.color": INK,
            "xtick.color": INK,
            "ytick.color": MUTED,
        }
    )
    fig, ax = plt.subplots(figsize=(8.6, 4.45), dpi=220)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    x = np.arange(len(order))
    width = 0.23

    for i, (col, color) in enumerate(series):
        vals = d.loc[order, col].to_numpy(dtype=float)
        bars = ax.bar(
            x + (i - 1) * width,
            vals,
            width=width,
            color=color,
            edgecolor="none",
            label=label_map[col],
            zorder=3,
        )
        for bar, val in zip(bars, vals):
            if val > 0:
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    val + 0.5,
                    f"{int(val)}",
                    ha="center",
                    va="bottom",
                    fontsize=7.6,
                    color=INK,
                )

    ax.set_xticks(x)
    ax.set_xticklabels(order, fontsize=9.5, fontweight="bold")
    ax.set_ylim(0, 27.5)
    ax.set_ylabel("Member-days above threshold", fontsize=9.5, labelpad=7)
    ax.yaxis.grid(True, color="#ECE7DC", linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color(LINE)
    ax.spines["bottom"].set_color(LINE)
    ax.tick_params(axis="both", length=0, pad=5)
    ax.legend(
        loc="upper right",
        frameon=False,
        fontsize=8.5,
        handlelength=1.2,
        labelspacing=0.35,
    )
    ax.set_title(
        "16 July 2025 initialization · 51 members · days 1–10",
        loc="left",
        fontsize=9,
        color=MUTED,
        pad=9,
    )
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------
def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = c(NAVY)

    add_rect(
        slide,
        0,
        0,
        13.333,
        0.10,
        fill=AMBER,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_rect(
        slide,
        0,
        7.40,
        13.333,
        0.10,
        fill=TEAL,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_line(slide, 8.92, 1.15, 8.92, 6.35, color="#2A4760", width=0.75)

    add_text(
        slide,
        0.78,
        0.92,
        7.6,
        0.30,
        "SUDAN  ·  JJA 2025",
        size=11,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.74,
        1.42,
        7.95,
        1.92,
        "Did the forecast warn\nof extreme rain?",
        size=42,
        color=WHITE,
        bold=True,
        line_spacing=0.92,
    )
    add_text(
        slide,
        0.78,
        3.50,
        7.45,
        1.30,
        "Local thresholds, ensemble signals, and the when-and-where of risk in five Sudanese towns.",
        size=17,
        color="#D9E2E8",
        line_spacing=1.08,
    )
    add_text(
        slide,
        0.78,
        5.02,
        7.6,
        0.34,
        "AFRICAN SUMMER SCHOOL ON AI FOR WEATHER AND CLIMATE MODELLING 2026",
        size=8.8,
        color="#9FB1BF",
        bold=True,
    )
    add_text(
        slide,
        0.78,
        5.42,
        6.0,
        0.28,
        "Group project · October 2026",
        size=11,
        color="#C7D2DC",
    )

    add_text(
        slide,
        9.32,
        1.25,
        3.10,
        0.26,
        "FIVE TOWNS",
        size=9.5,
        color=AMBER,
        bold=True,
    )
    towns = ["Khartoum", "Kassala", "Gedaref", "El Obeid", "Nyala"]
    for i, town in enumerate(towns):
        yy = 1.82 + i * 0.66
        add_text(
            slide,
            9.32,
            yy,
            0.38,
            0.24,
            f"0{i + 1}",
            size=8.5,
            color="#6F879A",
            bold=True,
        )
        add_text(
            slide,
            9.75,
            yy - 0.02,
            2.65,
            0.28,
            town.upper(),
            size=12,
            color=WHITE,
            bold=True,
        )
        if i < 4:
            add_line(
                slide, 9.32, yy + 0.39, 12.20, yy + 0.39, color="#234058", width=0.5
            )

    add_text(
        slide,
        9.32,
        5.58,
        2.95,
        0.52,
        "24.25–37.0°E\n11.5–16.25°N",
        size=9.5,
        color="#9FB1BF",
        line_spacing=1.2,
    )
    add_text(
        slide,
        0.78,
        6.52,
        5.6,
        0.26,
        "ECMWF IFS 49r1  ·  CMORPH 1998–2024",
        size=9,
        color="#7F95A7",
    )

    slide_notes(
        slide,
        "Opening question: did the JJA 2025 ECMWF ensemble produce locally meaningful "
        "signals of extreme rainfall in Sudan, and when and where did they appear? "
        "The deck follows five towns through Parts A to D, then closes with limits, "
        "reflections and the repository tag.",
    )
    return slide


def slide_region(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "01 / The place", "Five towns across Sudan's rainfall gradient")

    add_picture_fit(
        slide,
        ROOT / "out/rainfall_thresholds/figures/regional_10yr_24hr.png",
        0.65,
        1.55,
        7.05,
        3.48,
        alt="Map of the Sudan study region with five towns and the CMORPH 24-hour 10-year rainfall level.",
    )
    add_caption(
        slide,
        0.68,
        5.08,
        6.95,
        "Background: CMORPH 24-hour 10-year level; dots mark the nearest CMORPH pixel to each town.",
    )

    add_rect(slide, 7.95, 1.55, 4.68, 3.60, fill=CARD, line=LINE, radius=0.04)
    add_text(
        slide,
        8.17,
        1.73,
        2.9,
        0.24,
        "24-HOUR 2-YEAR LEVEL",
        size=8.5,
        color=AMBER_TEXT,
        bold=True,
    )
    town_data = [
        ("Khartoum", "23.8", "capital · Nile confluence"),
        ("Kassala", "37.9", "eastern border · Gash"),
        ("Gedaref", "53.0", "rain-fed agriculture"),
        ("El Obeid", "37.2", "central savanna"),
        ("Nyala", "32.3", "western Darfur"),
    ]
    for i, (town, value, description) in enumerate(town_data):
        yy = 2.10 + i * 0.56
        add_text(
            slide,
            8.17,
            yy,
            1.45,
            0.24,
            town,
            size=10.8,
            color=INK,
            bold=True,
        )
        add_text(
            slide,
            9.61,
            yy - 0.03,
            0.90,
            0.28,
            value,
            size=13.5,
            color=AMBER_TEXT,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
        add_text(
            slide,
            10.53,
            yy + 0.02,
            1.75,
            0.24,
            "mm/day",
            size=8.2,
            color=MUTED,
        )
        add_text(
            slide,
            8.17,
            yy + 0.25,
            2.30,
            0.20,
            description,
            size=8.2,
            color=MUTED,
        )
        if i < 4:
            add_line(slide, 8.17, yy + 0.48, 12.40, yy + 0.48, color="#EEE9DE", width=0.5)

    add_rect(slide, 0.65, 5.50, 12.0, 1.26, fill=NAVY, line=None, radius=0.035)
    add_text(
        slide,
        0.89,
        5.66,
        2.25,
        0.28,
        "WHY THESE TOWNS?",
        size=8.6,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.89,
        5.97,
        3.50,
        0.50,
        "The capital and four state capitals: people, markets, roads.",
        size=10.2,
        color=WHITE,
    )
    add_text(
        slide,
        4.50,
        5.97,
        3.50,
        0.50,
        "Rain-fed agriculture and local flood exposure.",
        size=10.2,
        color=WHITE,
    )
    add_text(
        slide,
        8.12,
        5.97,
        4.22,
        0.50,
        "A 2.2× spread in the local 2-year level: 23.8–53.0 mm/day.",
        size=10.2,
        color=WHITE,
    )
    add_footer(slide, 2)
    slide_notes(
        slide,
        "The domain is 24.25–37°E and 11.5–16.25°N. The five towns are the national "
        "capital and four state capitals. They span the rainfall gradient from arid "
        "Khartoum to the higher local thresholds in Gedaref. The map background is "
        "the 10-year level, which is a reference map; the numbers on the right are the "
        "2-year levels used in Parts B–D.",
    )
    return slide


def slide_data(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "02 / Data", "Two stores, one comparison frame")

    # Forecast card
    add_rect(slide, 0.65, 1.55, 5.95, 4.25, fill=CARD, line=LINE, radius=0.035)
    add_rect(
        slide,
        0.65,
        1.55,
        0.08,
        4.25,
        fill=AMBER,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_text(
        slide,
        0.93,
        1.72,
        4.8,
        0.24,
        "FORECAST",
        size=8.8,
        color=AMBER_TEXT,
        bold=True,
    )
    add_text(
        slide,
        0.93,
        1.97,
        4.9,
        0.36,
        "ECMWF IFS ensemble",
        size=18,
        color=NAVY,
        bold=True,
    )
    add_text(
        slide,
        0.93,
        2.34,
        4.95,
        0.24,
        "49r1 · JJA 2025 realized store",
        size=9.5,
        color=MUTED,
    )
    forecast_rows = [
        ("Members", "51"),
        ("Daily starts", "92  ·  1 Jun–31 Aug 2025"),
        ("Lead times", "0–240 h  ·  65 steps"),
        ("Grid", "0.25° (~28 km)  ·  321×261"),
        ("Variable", "tp cumulative, m → mm"),
        ("Snapshot", "212QY6MDW3RZ7PWD7A2G"),
    ]
    for i, (label, value) in enumerate(forecast_rows):
        yy = 2.79 + i * 0.48
        add_text(
            slide,
            0.93,
            yy,
            1.18,
            0.22,
            label.upper(),
            size=7.2,
            color=LIGHT_MUTED,
            bold=True,
        )
        add_text(
            slide,
            2.16,
            yy - 0.02,
            4.18,
            0.26,
            value,
            size=10.2 if label != "Snapshot" else 7.9,
            color=INK,
            bold=label == "Members",
            font=MONO if label == "Snapshot" else FONT,
        )

    # CMORPH card
    add_rect(slide, 6.78, 1.55, 5.90, 4.25, fill=CARD, line=LINE, radius=0.035)
    add_rect(
        slide,
        6.78,
        1.55,
        0.08,
        4.25,
        fill=TEAL,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_text(
        slide,
        7.06,
        1.72,
        4.8,
        0.24,
        "OBSERVED REFERENCE",
        size=8.8,
        color=TEAL,
        bold=True,
    )
    add_text(
        slide,
        7.06,
        1.97,
        4.9,
        0.36,
        "CMORPH return levels",
        size=18,
        color=NAVY,
        bold=True,
    )
    add_text(
        slide,
        7.06,
        2.34,
        4.95,
        0.24,
        "1998–2024 · 27 annual maxima",
        size=9.5,
        color=MUTED,
    )
    cmorph_rows = [
        ("Record", "1998–2024  ·  27 years"),
        ("Products", "9 durations  ×  6 return periods"),
        ("Grid", "~8 km  ·  550×474  ·  lat ascending"),
        ("Variable", "return_period_precip, mm"),
        ("Thresholds", "2 / 10 / 50-year  ·  24 / 48 h"),
        ("Snapshot", "309ZYV7EMDMNPD0EDZK0"),
    ]
    for i, (label, value) in enumerate(cmorph_rows):
        yy = 2.79 + i * 0.48
        add_text(
            slide,
            7.06,
            yy,
            1.05,
            0.22,
            label.upper(),
            size=7.2,
            color=LIGHT_MUTED,
            bold=True,
        )
        add_text(
            slide,
            8.18,
            yy - 0.02,
            4.30,
            0.26,
            value,
            size=10.2 if label != "Snapshot" else 7.9,
            color=INK,
            font=MONO if label == "Snapshot" else FONT,
        )

    add_rect(slide, 0.65, 6.00, 12.03, 0.94, fill=NAVY, line=None, radius=0.03)
    add_text(
        slide,
        0.88,
        6.10,
        3.2,
        0.22,
        "SHARED COMPARISON FRAME",
        size=8.3,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        6.34,
        11.5,
        0.24,
        "24-hour totals  ·  mm  ·  strict >  ·  nearest pixel  ·  local 2-year threshold for B/C/D",
        size=10.2,
        color=WHITE,
    )
    add_text(
        slide,
        0.88,
        6.64,
        11.6,
        0.20,
        "Source: data.source.coop/e4drr-project virtual stores. Only requested maps/steps are read.",
        size=7.5,
        color="#B9C7D2",
    )
    add_footer(slide, 3)
    slide_notes(
        slide,
        "The forecast store has 51 members and 92 daily starts from 1 June to 31 August "
        "2025, 0–240 hours, on a 0.25-degree latitude-longitude grid. The forecast "
        "variable is cumulative total precipitation in metres, so we difference whole-day "
        "leads and convert to millimetres. CMORPH provides return levels from 1998–2024, "
        "nine durations by six return periods, on a roughly 8 km grid. Both use strict "
        "greater-than and nearest-pixel matching.",
    )
    return slide


def slide_methods(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "03 / Method", "A common threshold, read four ways")

    cards = [
        (
            "A",
            "What is extreme?",
            "OBSERVED",
            "27 annual maxima → Gumbel fit → local 24-hour 2-year level. 10- and 50-year levels test sensitivity.",
            AMBER,
        ),
        (
            "B",
            "How much does the definition matter?",
            "DEFINITION",
            "16 Jul start: fixed 20 mm vs local 2-year vs 10-year. Count member-days above each threshold.",
            TEAL,
        ),
        (
            "C",
            "When did the signal appear?",
            "TIME",
            "21 daily starts, 1–21 Jul. Five towns; days 1–10; 51 members; count members above the local 2-year level.",
            NAVY,
        ),
        (
            "D",
            "Where was the risk highest?",
            "SPACE",
            "Every forecast grid point. Four stratified starts; days 1–5; probability and season-frequency maps.",
            ORANGE,
        ),
    ]

    outputs = [
        "local threshold table",
        "member-day counts",
        "start-date plume",
        "probability + frequency maps",
    ]
    for i, (letter, question, eyebrow, body, accent) in enumerate(cards):
        x = 0.65 + i * 3.05
        add_rect(slide, x, 1.55, 2.85, 4.35, fill=CARD, line=LINE, radius=0.035)
        add_rect(
            slide,
            x,
            1.55,
            2.85,
            0.075,
            fill=accent,
            line=None,
            radius=0.0,
            shape=MSO_SHAPE.RECTANGLE,
        )
        add_rect(
            slide,
            x + 0.20,
            1.80,
            0.52,
            0.52,
            fill=accent,
            line=None,
            radius=0.5,
            shape=MSO_SHAPE.OVAL,
        )
        add_text(
            slide,
            x + 0.20,
            1.84,
            0.52,
            0.42,
            letter,
            size=18,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            x + 0.84,
            1.82,
            1.85,
            0.22,
            eyebrow,
            size=7.8,
            color=accent,
            bold=True,
        )
        add_text(
            slide,
            x + 0.84,
            2.05,
            1.85,
            0.60,
            question,
            size=11.5,
            color=NAVY,
            bold=True,
            line_spacing=0.98,
        )
        add_text(
            slide,
            x + 0.20,
            2.85,
            2.45,
            1.95,
            body,
            size=9.8,
            color=INK,
            line_spacing=1.08,
        )
        add_line(slide, x + 0.20, 4.38, x + 2.65, 4.38, color=LINE, width=0.5)
        add_text(
            slide,
            x + 0.20,
            4.55,
            2.42,
            0.20,
            "OUTPUT",
            size=7.3,
            color=LIGHT_MUTED,
            bold=True,
        )
        add_text(
            slide,
            x + 0.20,
            4.74,
            2.42,
            0.42,
            outputs[i],
            size=8.7,
            color=MUTED,
            line_spacing=0.98,
        )
        if letter == "A":
            add_text(
                slide,
                x + 0.20,
                5.34,
                2.42,
                0.40,
                "Feeds B, C and D",
                size=8.4,
                color=AMBER_TEXT,
                bold=True,
            )
        else:
            add_text(
                slide,
                x + 0.20,
                5.34,
                2.42,
                0.40,
                "Uses the A threshold",
                size=8.4,
                color=LIGHT_MUTED,
            )

    add_rect(slide, 0.65, 6.13, 12.03, 0.56, fill=AMBER_LIGHT, line=None, radius=0.05)
    add_text(
        slide,
        0.88,
        6.23,
        11.6,
        0.26,
        "A sets the reference; B tests the definition; C adds time; D adds space.",
        size=10.8,
        color=INK,
        bold=True,
    )
    add_footer(slide, 4)
    slide_notes(
        slide,
        "The four parts share one reference threshold: the local 24-hour 2-year level "
        "from CMORPH. A estimates that level. B asks whether the answer changes when the "
        "threshold changes. C holds the threshold fixed and adds forecast start date and "
        "lead time. D holds the threshold fixed and applies it at every forecast grid point. "
        "This is why the parts are comparable rather than four disconnected exercises.",
    )
    return slide


def slide_lineage(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "04 / Lineage", "Borrowed patterns, rewritten for Sudan")

    add_rect(slide, 0.65, 1.55, 3.32, 3.85, fill=NAVY, line=None, radius=0.035)
    add_text(
        slide,
        0.88,
        1.74,
        2.4,
        0.24,
        "WORKSHOP SOURCE",
        size=8.6,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        2.02,
        2.8,
        0.55,
        "E1–E8",
        size=28,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        2.62,
        2.6,
        0.30,
        "Day 2 · python-workshop",
        size=9.5,
        color="#C7D2DC",
    )
    add_bullets(
        slide,
        0.88,
        3.06,
        2.65,
        1.85,
        [
            "Examples were read-only inspiration.",
            "Analysis scripts live in this repo's scripts/ folder.",
            "Every figure comes from a committed script.",
        ],
        size=9.6,
        color=WHITE,
        bullet_color=AMBER,
        space_after=6,
    )
    add_text(
        slide,
        0.88,
        5.05,
        2.55,
        0.25,
        "Adapted, not copied.",
        size=9.2,
        color=AMBER,
        bold=True,
    )

    table_x = 4.22
    table_y = 1.55
    col_widths = [0.92, 4.18, 3.00]
    row_h = 0.535
    header_h = 0.45
    rows = [
        ("E1–E2", "Open a virtual store; read a realized plume", "Store access; member reads"),
        ("E3", "Pan vs pencil chunk reads", "D: crop region/steps before compute"),
        ("E4", "CMORPH return periods and Gumbel fit", "A: local thresholds"),
        ("E5–E6", "xclim heavy days; forecast vs local extremes", "B/C: member counts"),
        ("E7", "Review before push", "Git checks and reproducibility"),
        ("E8", "Local-threshold probability map", "D: regional risk maps"),
    ]
    # Header
    x = table_x
    for width, label in zip(col_widths, ["EXERCISE", "WHAT WE BORROWED", "WHERE IT LANDED"]):
        add_rect(
            slide,
            x,
            table_y,
            width,
            header_h,
            fill=NAVY,
            line=None,
            radius=0.0,
            shape=MSO_SHAPE.RECTANGLE,
        )
        add_text(
            slide,
            x + 0.10,
            table_y + 0.13,
            width - 0.18,
            0.20,
            label,
            size=7.7,
            color=WHITE,
            bold=True,
        )
        x += width

    for i, (ex, borrowed, landed) in enumerate(rows):
        y = table_y + header_h + i * row_h
        fill = CARD if i % 2 == 0 else "#F4F0E7"
        x = table_x
        for width in col_widths:
            add_rect(
                slide,
                x,
                y,
                width,
                row_h,
                fill=fill,
                line=None,
                radius=0.0,
                shape=MSO_SHAPE.RECTANGLE,
            )
            x += width
        add_line(slide, table_x, y + row_h, table_x + sum(col_widths), y + row_h, color=LINE, width=0.5)
        add_text(
            slide,
            table_x + 0.10,
            y + 0.17,
            0.80,
            0.22,
            ex,
            size=9.3,
            color=AMBER_TEXT,
            bold=True,
        )
        add_text(
            slide,
            table_x + col_widths[0] + 0.10,
            y + 0.13,
            col_widths[1] - 0.18,
            0.38,
            borrowed,
            size=8.7,
            color=INK,
            line_spacing=0.95,
        )
        add_text(
            slide,
            table_x + col_widths[0] + col_widths[1] + 0.10,
            y + 0.13,
            col_widths[2] - 0.18,
            0.38,
            landed,
            size=8.7,
            color=MUTED,
            line_spacing=0.95,
        )

    add_rect(slide, 0.65, 5.58, 12.03, 0.82, fill=CARD, line=LINE, radius=0.03)
    add_text(
        slide,
        0.88,
        5.68,
        2.0,
        0.20,
        "REPO SCRIPTS",
        size=7.7,
        color=AMBER_TEXT,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        5.94,
        11.5,
        0.30,
        "scripts/rainfall_thresholds.py  ·  forecast_threshold_counts.py  ·  forecast_signal_leadtime.py  ·  region_risk_maps.py",
        size=8.0,
        color=INK,
        font=MONO,
    )
    add_text(
        slide,
        0.88,
        6.20,
        11.6,
        0.18,
        "Examples inspired the pattern; the Sudan scripts, settings, tests and outputs are ours.",
        size=7.8,
        color=MUTED,
    )
    add_footer(slide, 5)
    slide_notes(
        slide,
        "The workshop provided eight patterns, but not a ready-made Sudan analysis. "
        "We adapted the store-opening and plume pattern from E1/E2, the chunk-read lesson "
        "from E3, the return-period logic from E4, the heavy-day and local-threshold counts "
        "from E5/E6, the review discipline from E7 and the probability map from E8. The "
        "analysis scripts and settings are in this repository, not copied from the workshop.",
    )
    return slide


def slide_part_a(prs, asset_dir):
    chart = asset_dir / "part_a_thresholds.png"
    make_part_a_chart(chart)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "05 / Finding A", "What counts as extreme? It depends where you are.")

    add_picture_fit(
        slide,
        chart,
        0.65,
        1.55,
        7.45,
        3.95,
        border=False,
        alt="Grouped bar chart of 24-hour 2-, 10- and 50-year rainfall thresholds for the five towns.",
    )
    add_caption(
        slide,
        0.68,
        5.55,
        7.4,
        "24-hour levels, mm. The 2-year level is the primary threshold used in B–D.",
    )

    add_text(
        slide,
        8.38,
        1.58,
        2.2,
        0.24,
        "24-HOUR 2-YEAR LEVEL",
        size=8.3,
        color=AMBER_TEXT,
        bold=True,
    )
    add_text(
        slide,
        8.35,
        1.82,
        4.1,
        0.55,
        "23.8–53.0",
        size=30,
        color=NAVY,
        bold=True,
    )
    add_text(
        slide,
        8.40,
        2.38,
        3.8,
        0.24,
        "mm/day across the five towns",
        size=10.2,
        color=MUTED,
    )
    add_bullets(
        slide,
        8.38,
        2.78,
        4.05,
        2.52,
        [
            ("Gedaref is highest; ", "Khartoum is lowest at the 2-year level."),
            ("50-year levels: ", "65.8–90.3 mm/day."),
            ("Fit sensitivity: ", "removing the largest year changes the 50-year fit by 3–9%."),
        ],
        size=10.9,
        space_after=7,
    )
    add_rect(slide, 8.38, 5.42, 4.05, 0.86, fill=AMBER_LIGHT, line=None, radius=0.04)
    add_text(
        slide,
        8.55,
        5.53,
        3.70,
        0.64,
        "A fixed national threshold would misread both ends of this gradient.",
        size=10.2,
        color=INK,
        bold=True,
        line_spacing=1.0,
    )
    add_footer(slide, 6)
    slide_notes(
        slide,
        "Part A fits a Gumbel distribution to 27 annual maxima at each town's nearest "
        "CMORPH pixel. The 24-hour 2-year level ranges from 23.8 mm/day in Khartoum to "
        "53.0 mm/day in Gedaref, a factor of 2.2. The 10- and 50-year levels are also shown. "
        "A fixed national threshold would therefore label very different rainfall amounts "
        "as equally extreme.",
    )
    return slide


def slide_part_b(prs, asset_dir):
    chart = asset_dir / "part_b_thresholds.png"
    make_part_b_chart(chart)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "06 / Finding B", "How much does the threshold definition matter?")

    add_picture_fit(
        slide,
        chart,
        0.65,
        1.55,
        7.45,
        3.95,
        border=False,
        alt="Grouped bar chart of member-days above fixed 20 mm, local 2-year and 10-year thresholds.",
    )
    add_caption(
        slide,
        0.68,
        5.55,
        7.4,
        "One 16 July 2025 initialization; member-days summed over forecast days 1–10.",
    )

    add_text(
        slide,
        8.38,
        1.58,
        4.0,
        0.24,
        "MEMBER-DAYS ABOVE THRESHOLD",
        size=8.3,
        color=AMBER_TEXT,
        bold=True,
    )
    add_text(
        slide,
        8.35,
        1.82,
        4.15,
        0.55,
        "46 → 3 → 0",
        size=30,
        color=NAVY,
        bold=True,
    )
    add_text(
        slide,
        8.40,
        2.38,
        3.90,
        0.24,
        "fixed 20 mm  →  2-year  →  10-year",
        size=9.8,
        color=MUTED,
    )
    add_bullets(
        slide,
        8.38,
        2.78,
        4.05,
        2.65,
        [
            ("Fixed 20 mm: ", "46 member-days; Gedaref alone contributes 24."),
            ("Local 2-year: ", "3 member-days (Nyala 2, El Obeid 1)."),
            ("Local 10-year: ", "none."),
            ("The ranking changes: ", "the definition decides whether a warning exists at all."),
        ],
        size=10.4,
        space_after=5,
    )
    add_rect(slide, 8.38, 5.60, 4.05, 0.68, fill=TEAL_LIGHT, line=None, radius=0.04)
    add_text(
        slide,
        8.54,
        5.72,
        3.72,
        0.44,
        "Same forecast, same day — only the threshold changed.",
        size=9.7,
        color=INK,
        bold=True,
    )
    add_footer(slide, 7)
    slide_notes(
        slide,
        "Part B holds the forecast fixed at the 16 July initialization and changes only "
        "the threshold. A fixed 20 mm/day threshold produces 46 member-days above the "
        "threshold. The local 2-year level produces only three, and the 10-year level "
        "produces none. The ranking also changes: Gedaref dominates the fixed threshold, "
        "while Nyala and El Obeid are the only towns with members above their local "
        "2-year levels. The definition is not cosmetic.",
    )
    return slide


def slide_part_c(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "07 / Finding C", "When did the signal appear?")

    add_picture_fit(
        slide,
        ROOT / "out/forecast_leadtime/08_forecast_leadtime_plume_nyala_20250707.png",
        0.65,
        1.55,
        7.55,
        3.00,
        alt="Nyala daily rainfall plume for the 7 July 2025 forecast start, with the 2-year threshold and member exceedance counts.",
    )
    add_caption(
        slide,
        0.68,
        4.55,
        7.45,
        "Nyala · 7 July start · red dashed line = 32.3 mm 2-year level · bars = members above.",
    )
    add_rect(
        slide,
        0.65,
        4.98,
        0.055,
        0.38,
        fill=AMBER,
        line=None,
        radius=0.0,
        shape=MSO_SHAPE.RECTANGLE,
    )
    add_text(
        slide,
        0.82,
        4.96,
        7.30,
        0.42,
        "Early signal, but the strongest and most consistent was Nyala.",
        size=12.2,
        color=NAVY,
        bold=True,
    )

    add_text(
        slide,
        8.38,
        1.58,
        3.8,
        0.24,
        "MOST CONSISTENT SIGNAL",
        size=8.3,
        color=AMBER_TEXT,
        bold=True,
    )
    add_text(
        slide,
        8.35,
        1.82,
        3.8,
        0.55,
        "20 / 21",
        size=30,
        color=NAVY,
        bold=True,
    )
    add_text(
        slide,
        8.40,
        2.38,
        3.75,
        0.42,
        "Nyala starts with ≥1 member above its 32.3 mm local level.",
        size=9.2,
        color=MUTED,
        line_spacing=1.02,
    )
    add_bullets(
        slide,
        8.38,
        2.92,
        4.05,
        2.30,
        [
            ("72 member-days, ", "peak 4 members in one day."),
            ("Earliest valid date: ", "Kassala, 4 July — from the 2 July start, day 2."),
            ("Most alarming start: ", "Nyala 7 July — 10 member-days, peak 4 on 14 July."),
            ("Contrast: ", "Khartoum 18/21 starts; Gedaref 3/21."),
        ],
        size=10.2,
        space_after=4,
    )
    add_rect(slide, 0.65, 5.55, 12.03, 0.78, fill=NAVY, line=None, radius=0.035)
    add_text(
        slide,
        0.88,
        5.64,
        3.0,
        0.20,
        "WHEN TO LOOK",
        size=7.8,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        5.89,
        11.4,
        0.26,
        "4 Jul  earliest valid signal (Kassala)   ·   14 Jul  peak (Nyala)   ·   1–21 Jul  start window",
        size=9.6,
        color=WHITE,
    )
    add_text(
        slide,
        8.38,
        5.22,
        3.9,
        0.22,
        "Member counts, not probabilities.",
        size=8.2,
        color=MUTED,
        italic=True,
    )
    add_footer(slide, 8)
    slide_notes(
        slide,
        "Part C fixes the local 2-year threshold and looks across 21 daily starts from "
        "1–21 July. Nyala is the most consistent and strongest signal: 20 of 21 starts have "
        "at least one member above their local level, with 72 member-days and a peak of four members in a day. "
        "Kassala provides the earliest valid calendar date, 4 July, from the 2 July start at "
        "day 2. The most alarming individual start is Nyala on 7 July: 10 member-days, "
        "peaking at four members on 14 July. These are member counts, not probabilities.",
    )
    return slide


def slide_part_d(prs, asset_dir):
    # Crop the day-5 panel and colour bar from the five-panel probability figure.
    source = Image.open(ROOT / "out/region_risk/09_probability_map_20250819_p20.png").convert("RGB")
    crop_path = asset_dir / "part_d_day5_probability.png"
    source.crop((1720, 60, 2322, 560)).save(crop_path)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "08 / Finding D", "Where was the risk concentrated?")

    add_picture_fit(
        slide,
        ROOT / "out/region_risk/09_frequency_map_p20.png",
        0.65,
        1.55,
        7.55,
        3.58,
        alt="Season frequency map of start-day cases with at least 20 percent probability; highest values are in the west near Nyala.",
    )
    add_picture_fit(
        slide,
        crop_path,
        8.45,
        1.55,
        4.18,
        3.58,
        alt="Probability map for 19 August 2025 day 5; the largest probability is about 39 percent near 12.75 north, 24.25 east.",
    )
    add_caption(
        slide,
        0.68,
        5.17,
        7.45,
        "20 cases = 4 starts (3 Jun, 11 Jul, 1 Aug, 19 Aug) × 5 days.\nColour: number of cases with probability ≥20%; maximum 1.",
    )
    add_caption(
        slide,
        8.48,
        5.17,
        4.12,
        "19 Aug 2025 · day 5 (24 Aug) · black contour = 20%.",
    )

    stats = [
        ("39.2%", "Peak probability", "12.75°N, 24.25°E\n19 Aug · day 5"),
        ("18", "Grid cells ≥20%", "at least once\nmax frequency 1 of 20"),
        ("5 / 20", "At a 10% threshold", "max frequency\nsame western focus"),
    ]
    for i, (value, label, detail) in enumerate(stats):
        x = 0.65 + i * 4.06
        add_rect(slide, x, 5.72, 3.85, 1.05, fill=CARD, line=LINE, radius=0.035)
        add_text(
            slide,
            x + 0.18,
            5.80,
            1.20,
            0.40,
            value,
            size=19,
            color=AMBER_TEXT if i == 0 else NAVY,
            bold=True,
        )
        add_text(
            slide,
            x + 1.36,
            5.82,
            2.28,
            0.22,
            label,
            size=8.8,
            color=INK,
            bold=True,
        )
        add_text(
            slide,
            x + 1.36,
            6.06,
            2.28,
            0.52,
            detail,
            size=7.7,
            color=MUTED,
            line_spacing=0.98,
        )
    add_footer(slide, 9)
    slide_notes(
        slide,
        "Part D applies the local 2-year threshold at every forecast grid point. Four "
        "stratified starts sample the season: 3 June, 11 July, 1 August and 19 August, "
        "with days 1–5. The strongest cell is 39.2% at 12.75°N, 24.25°E on day 5 of the "
        "19 August start, valid 24 August. Eighteen cells reached 20% at least once, but "
        "no cell did so in more than 1 of 20 cases. The west/south-west near Nyala dominates; "
        "the eastern towns stay quiet. A 10% sensitivity run raises the maximum frequency "
        "to 5 of 20 cases at 13.0°N, 24.25°E.",
    )
    return slide


def slide_limits(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "09 / Limits", "What this analysis cannot tell us")

    cards = [
        (
            "No 2025 verification",
            "Observed 2025 rainfall is not in these stores. We describe forecast signals, not confirmed events.",
            "Signals, not confirmed events",
            AMBER,
        ),
        (
            "Short record, extrapolation",
            "27 annual maxima. 50-year levels are model estimates; removing the largest year moves them 3–9%.",
            "Treat 50-year values as indicative",
            TEAL,
        ),
        (
            "Scale mismatch",
            "Forecast ~28 km vs CMORPH ~8 km. Nearest pixels do not represent the same area; satellite estimates differ from gauges.",
            "Compare at compatible scales",
            NAVY,
        ),
        (
            "A threshold is not an impact",
            "The 2-year annual-max level is not a crop, flood or infrastructure damage threshold. Only 24-hour totals are used.",
            "Pair with impact thresholds",
            ORANGE,
        ),
        (
            "Sampling limits",
            "Part C covers July only. Part D uses four starts × five days; days within one start are correlated.",
            "Do not extend to the full season",
            TEAL,
        ),
        (
            "Counts are not probabilities",
            "Part C counts members; Part D uses the fraction of 51 members. Neither is calibrated or verified here.",
            "Do not read as probability of impact",
            AMBER,
        ),
    ]
    for i, (title, body, implication, accent) in enumerate(cards):
        row = i // 3
        col = i % 3
        x = 0.65 + col * 4.06
        y = 1.55 + row * 2.35
        add_rect(slide, x, y, 3.85, 2.15, fill=CARD, line=LINE, radius=0.035)
        add_rect(
            slide,
            x,
            y,
            0.075,
            2.15,
            fill=accent,
            line=None,
            radius=0.0,
            shape=MSO_SHAPE.RECTANGLE,
        )
        add_text(
            slide,
            x + 0.22,
            y + 0.18,
            0.35,
            0.22,
            f"0{i + 1}",
            size=8,
            color=accent,
            bold=True,
        )
        add_text(
            slide,
            x + 0.60,
            y + 0.16,
            3.02,
            0.44,
            title,
            size=11.0,
            color=NAVY,
            bold=True,
            line_spacing=0.95,
        )
        add_text(
            slide,
            x + 0.22,
            y + 0.76,
            3.38,
            1.18,
            body,
            size=9.5,
            color=INK,
            line_spacing=1.07,
        )
        add_line(slide, x + 0.22, y + 1.50, x + 3.56, y + 1.50, color="#EEE9DE", width=0.5)
        add_text(
            slide,
            x + 0.22,
            y + 1.62,
            3.30,
            0.28,
            "→  " + implication,
            size=8.2,
            color=accent,
            bold=True,
        )

    add_rect(slide, 0.65, 6.18, 12.03, 0.52, fill=NAVY, line=None, radius=0.03)
    add_text(
        slide,
        0.88,
        6.28,
        11.60,
        0.25,
        "Next: verify against 2025 observations · calibrate thresholds · extend the sampling window.",
        size=10.2,
        color=WHITE,
        bold=True,
    )
    add_footer(slide, 10)
    slide_notes(
        slide,
        "The most important limitation is that the 2025 observations are absent from these "
        "stores. We can describe the forecast signal and its timing, but not verify that "
        "the events happened. The 27-year record makes 50-year extrapolation uncertain; "
        "the forecast and CMORPH grids differ by a factor of about three; the 2-year "
        "annual-maximum level is not an impact threshold; and the Part D sample is four "
        "starts by five days, not an exhaustive season.",
    )
    return slide


def slide_reflections(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c(PAPER)
    add_header(slide, "10 / Reflection", "From exercise scripts to a reusable workflow")

    columns = [
        (
            "WHAT WE ADAPTED",
            AMBER,
            [
                ("E1/E2 → ", "store access and realized plumes."),
                ("E3 → ", "pan-chunk selection before compute."),
                ("E4 → ", "Gumbel return levels; E5/E6 → member counts."),
                ("E7 → ", "review and version before push."),
                ("E8 → ", "local-threshold probability maps."),
            ],
        ),
        (
            "WHERE THE AGENT HELPED",
            TEAL,
            [
                ("Scaffolding: ", "turned prompts into runnable scripts."),
                ("Refactoring: ", "shared settings, helpers and tests."),
                ("Iteration: ", "plotting, checks and docs."),
                ("Speed: ", "repeated runs and output regeneration."),
                ("Checks: ", "tests, lint and output diffs."),
            ],
        ),
        (
            "WHAT WE CHECKED / CORRECTED",
            ORANGE,
            [
                ("Units: ", "metres vs millimetres; cumulative vs daily totals."),
                ("Geometry: ", "latitude order, nearest pixels, coverage."),
                ("Validity: ", "all-NaN starts; strict > semantics."),
                ("Interpretation: ", "counts ≠ probabilities; no 2025 verification."),
                ("Reproducibility: ", "fixed settings and committed outputs."),
            ],
        ),
    ]
    for i, (heading, accent, bullets) in enumerate(columns):
        x = 0.65 + i * 4.06
        add_rect(slide, x, 1.55, 3.85, 2.95, fill=CARD, line=LINE, radius=0.035)
        add_rect(
            slide,
            x,
            1.55,
            3.85,
            0.075,
            fill=accent,
            line=None,
            radius=0.0,
            shape=MSO_SHAPE.RECTANGLE,
        )
        add_text(
            slide,
            x + 0.22,
            1.75,
            3.35,
            0.26,
            heading,
            size=8.2,
            color=accent,
            bold=True,
        )
        add_bullets(
            slide,
            x + 0.22,
            2.14,
            3.35,
            2.15,
            bullets,
            size=9.7,
            bullet_color=accent,
            space_after=7,
            line_spacing=1.02,
        )

    # Compact workflow strip: the pattern we would reuse next time.
    add_text(
        slide,
        0.65,
        4.82,
        1.0,
        0.20,
        "WORKFLOW",
        size=7.6,
        color=AMBER_TEXT,
        bold=True,
    )
    steps = ["Prompt", "Script", "Test", "Figure", "Review", "Commit"]
    for i, step in enumerate(steps):
        x = 1.65 + i * 1.77
        add_rect(slide, x, 4.72, 1.58, 0.42, fill=CARD, line=LINE, radius=0.18)
        add_text(
            slide,
            x,
            4.79,
            1.58,
            0.24,
            step,
            size=8.4,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if i < len(steps) - 1:
            add_text(
                slide,
                x + 1.60,
                4.79,
                0.16,
                0.24,
                "→",
                size=9,
                color=LIGHT_MUTED,
                align=PP_ALIGN.CENTER,
            )

    add_rect(slide, 0.65, 5.45, 12.03, 0.88, fill=NAVY, line=None, radius=0.035)
    add_text(
        slide,
        0.88,
        5.55,
        8.15,
        0.62,
        "Day-to-day: reuse this pipeline for any ensemble hazard — local threshold, member counts, when/where maps, reviewed outputs.",
        size=9.7,
        color=WHITE,
        line_spacing=1.0,
    )
    add_text(
        slide,
        9.18,
        5.55,
        3.25,
        0.62,
        "The agent accelerates;\nthe analyst verifies.",
        size=9.9,
        color=AMBER,
        bold=True,
        line_spacing=1.0,
    )
    add_footer(slide, 11)
    slide_notes(
        slide,
        "The workshop exercises were a starting vocabulary, not a finished analysis. "
        "We adapted the patterns into scripts with shared settings and tests. The agent was "
        "most useful for scaffolding, refactoring, plotting and repeated runs. We had to "
        "check units, coordinate order, coverage, all-NaN starts and the strict greater-than "
        "semantics, and we had to keep counts separate from probabilities. For day-to-day "
        "work, the same pipeline generalises to other ensemble hazards: define a local "
        "threshold, count members, map when and where, and version the reviewed outputs.",
    )
    return slide


def slide_repo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = c(NAVY)
    add_header(slide, "11 / Reproduce", "Open, reviewed, reproducible", dark=True)

    repo_url = "https://github.com/BrianMburu/acmad-eac-summer-school"
    tag_url = "https://github.com/BrianMburu/acmad-eac-summer-school/tree/final"

    add_text(
        slide,
        0.65,
        1.72,
        8.0,
        0.26,
        "REPOSITORY",
        size=8.6,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.65,
        1.98,
        8.2,
        0.40,
        "github.com/BrianMburu/acmad-eac-summer-school",
        size=17,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        7.32,
        2.02,
        0.60,
        0.28,
        "OPEN",
        size=8.2,
        color=AMBER,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_link_overlay(slide, 0.65, 1.95, 7.25, 0.46, repo_url)
    add_line(slide, 0.65, 2.52, 8.30, 2.52, color="#2A4760", width=0.75)

    add_text(
        slide,
        0.65,
        2.72,
        2.0,
        0.24,
        "FINAL TAG",
        size=8.6,
        color=AMBER,
        bold=True,
    )
    add_chip(slide, 0.65, 3.00, 1.20, 0.48, "final", fill=AMBER, color=NAVY, size=17)
    add_text(
        slide,
        2.03,
        3.06,
        5.6,
        0.34,
        "annotated tag · deck + analysis outputs",
        size=9.8,
        color="#C7D2DC",
    )
    add_link_overlay(slide, 0.65, 3.00, 1.20, 0.48, tag_url)
    add_link_overlay(slide, 2.03, 3.02, 5.60, 0.40, tag_url)
    add_text(
        slide,
        0.65,
        3.72,
        7.75,
        0.52,
        "The tag points to the presentation, build script, analysis scripts, figures and CSV outputs.",
        size=9.9,
        color="#B9C7D2",
        line_spacing=1.08,
    )

    add_text(
        slide,
        0.65,
        4.62,
        8.1,
        0.24,
        "BUILD AND CHECK",
        size=8.2,
        color=AMBER,
        bold=True,
    )
    add_text(
        slide,
        0.65,
        4.88,
        8.2,
        0.68,
        "uv run --with python-pptx python presentation/build_deck.py\nuv run pytest   ·   uvx ruff check scripts tests presentation",
        size=8.1,
        color="#D9E2E8",
        font=MONO,
        line_spacing=1.18,
    )

    add_rect(slide, 9.05, 1.72, 3.60, 3.60, fill=NAVY_2, line="#2A4760", radius=0.035)
    add_text(
        slide,
        9.28,
        1.90,
        3.0,
        0.24,
        "AT A GLANCE",
        size=8.2,
        color=AMBER,
        bold=True,
    )
    stats = [
        ("4", "analyses"),
        ("5", "towns"),
        ("27", "years observed"),
        ("96", "tests passing"),
    ]
    for i, (value, label) in enumerate(stats):
        col = i % 2
        row = i // 2
        x = 9.28 + col * 1.52
        y = 2.28 + row * 1.22
        add_text(
            slide,
            x,
            y,
            1.25,
            0.44,
            value,
            size=23,
            color=WHITE,
            bold=True,
        )
        add_text(
            slide,
            x,
            y + 0.48,
            1.25,
            0.24,
            label,
            size=8.1,
            color="#B9C7D2",
        )
    add_text(
        slide,
        9.28,
        4.72,
        3.02,
        0.34,
        "scripts/ · out/ · tests/",
        size=8.0,
        color="#B9C7D2",
        font=MONO,
    )

    add_rect(slide, 0.65, 5.58, 12.03, 0.92, fill="#102B43", line="#2A4760", radius=0.03)
    add_text(
        slide,
        0.88,
        5.70,
        11.2,
        0.30,
        "Local thresholds turn a generic forecast into a place-specific, time-specific warning.",
        size=13.2,
        color=WHITE,
        bold=True,
    )
    add_text(
        slide,
        0.88,
        6.08,
        11.2,
        0.24,
        "Questions / discussion  ·  Group project, ACMAD AI summer school 2026",
        size=9.4,
        color="#B9C7D2",
    )
    add_footer(slide, 12, dark=True)
    slide_notes(
        slide,
        "The repository is public. The annotated tag final points to this deck, the build "
        "script, the four analysis scripts, the figures and the CSV outputs. The README "
        "records the methods and numbers; the tests and lint checks are part of the "
        "handoff. Closing message: local thresholds are what turn a generic rainfall "
        "forecast into a place-specific and time-specific warning.",
    )
    return slide


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="acmad-deck-assets-") as tmp:
        asset_dir = Path(tmp)
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        prs.core_properties.title = "Did the forecast warn of extreme rain? Sudan JJA 2025"
        prs.core_properties.subject = "Local rainfall thresholds and ECMWF ensemble signals for five Sudanese towns"
        prs.core_properties.author = "Sudan group, ACMAD AI summer school 2026"
        prs.core_properties.comments = "Built from committed figures and CSVs in the acmad-eac-summer-school repository."

        slide_title(prs)
        slide_region(prs)
        slide_data(prs)
        slide_methods(prs)
        slide_lineage(prs)
        slide_part_a(prs, asset_dir)
        slide_part_b(prs, asset_dir)
        slide_part_c(prs)
        slide_part_d(prs, asset_dir)
        slide_limits(prs)
        slide_reflections(prs)
        slide_repo(prs)

        prs.save(OUT_FILE)

    print(f"Wrote {OUT_FILE} ({OUT_FILE.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
